import io, json
from docx import Document
from docx.shared import Pt, RGBColor

def display(value):
    if value is None or value==[]:return 'ไม่ระบุ'
    if isinstance(value,list):return ' / '.join(map(str,value))
    return str(value)

def plain(job):
    if job.get('report_text') is not None:return job['report_text']
    if not job.get('report'):return ''
    r=json.loads(job['report']);meta=r['metadata']
    lines=[job['title'],'วันประชุม: '+display(meta.get('date')),'ผู้เข้าร่วม: '+display(meta.get('participants')),
           'ขอบเขต: '+display(r['coverage'].get('status')),'']
    def item(x):
        # Reports intentionally omit tape timestamps; the transcript keeps them for review.
        return '- '+x['text']
    for key,title in [('summary','ภาพรวมสำคัญ'),('topics','รายละเอียดแยกตามประเด็น'),('decisions','มติและข้อตกลง'),('proposals','ข้อเสนอที่ยังไม่เป็นมติ'),('actions','งานที่ต้องดำเนินการ'),('open_questions','ประเด็นค้างและคำถาม'),('review_notes','จุดที่ต้องตรวจทาน')]:
        lines+=['',title]
        if not r[key]:lines.append('ไม่พบข้อมูลที่ยืนยันในข้อความที่ได้รับ');continue
        for x in r[key]:
            if key=='topics':lines+=[x['title']]+[item(d) for d in x['details']]
            elif key=='actions':lines.append('- '+x['task']+' | ผู้รับผิดชอบ: '+display(x.get('owner'))+' | กำหนดส่ง: '+display(x.get('due_text'))+' | สถานะ: '+display(x.get('status_text'))+' | เงื่อนไข: '+display(x.get('conditions')))
            else:
                lines.append(item(x))
                if x.get('conditions'):lines.append('  เงื่อนไข: '+x['conditions'])
                if x.get('proposed_by'):lines.append('  ผู้เสนอ: '+x['proposed_by'])
    lines+=['','นัดหมายครั้งถัดไป',display(r['next_meeting'].get('time_text')),display(r['next_meeting'].get('purpose'))]
    return '\n'.join(lines)

def rtf_escape(text):
    # RTF with \uN? escapes keeps Thai intact for Word and LibreOffice without
    # relying on the ANSI code page of the viewer.
    out=[]
    for ch in text:
        code=ord(ch)
        if ch=='\\':out.append('\\\\')
        elif ch=='{':out.append('\\{')
        elif ch=='}':out.append('\\}')
        elif 32<=code<127:out.append(ch)
        elif code>0xffff:
            for unit in (0xd800+((code-0x10000)>>10),0xdc00+((code-0x10000)&0x3ff)):
                out.append('\\u%d?'%(unit-65536 if unit>32767 else unit))
        else:out.append('\\u%d?'%(code-65536 if code>32767 else code))
    return ''.join(out)

def doc_bytes(job,font):
    # Word-compatible RTF saved as .doc: no Word installation needed and the
    # legacy format stays readable in Word, LibreOffice and Pages.
    import re
    safe=font if isinstance(font,str) and re.fullmatch(r'[A-Za-z0-9 ]{1,64}',font) else 'Tahoma'
    status='ตรวจและอนุมัติแล้ว' if job['state']=='approved' else 'ฉบับร่างสำหรับตรวจทาน'
    parts=[f'{{\\rtf1\\ansi\\ansicpg874\\uc1\\deff0{{\\fonttbl{{\\f0\\fnil\\fcharset222 {safe};}}}}\\viewkind4\\fs24',
           '\\pard\\sa200\\b\\fs36 '+rtf_escape(job['title'])+'\\b0\\fs24\\par',
           '\\pard\\sa240 '+rtf_escape('สถานะ: '+status)+'\\par']
    for line in plain(job).splitlines():
        parts.append('\\pard\\sa120\\sl276\\slmult1 '+rtf_escape(line)+'\\par\n' if line else '\\pard\\par\n')
    parts.append('}')
    return ''.join(parts).encode('ascii')

def docx_bytes(job,font):
    d=Document();style=d.styles['Normal'];style.font.name=font;style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
    d.add_heading(job['title'],0)
    d.add_paragraph('สถานะ: '+('ตรวจและอนุมัติแล้ว' if job['state']=='approved' else 'ฉบับร่างสำหรับตรวจทาน'))
    for line in plain(job).splitlines():d.add_paragraph(line)
    for para in d.paragraphs:
        for run in para.runs:run.font.name=font;run.font.color.rgb=RGBColor(0,0,0)
    buf=io.BytesIO();d.save(buf);return buf.getvalue()
